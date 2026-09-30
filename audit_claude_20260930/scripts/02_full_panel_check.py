# -*- coding: utf-8 -*-
"""Full reconstruction check of verified_results/verified_master_panel.csv against the ORIGINAL KOSIS files.

Scope (what this script can and cannot show)
  * Numeric check: every crop-year value of area_ha, yield_kg10a, production_t, total_orchard_area_ha and
    bearing_area_ha in the verified panel is compared with the national cell of the original KOSIS file
    identified by (table_id, kosis_production_label_original).
  * The .xlsx exports are read with openpyxl (a different code path from the pandas reader used to build
    the panel). The SpreadsheetML (.xls) exports are read with the SAME xml2003 parser used in the build;
    for those files this is a re-run of the same parser, not an independent parser.
  * The script and the original build were written by the same author (Claude). It therefore verifies
    reproducibility from the raw files, not the suitability of the KOSIS definitions.
  * Treatment years: only an internal-consistency check against verified_results/verified_treatment_coding.csv.
    Original-source (Yearbook/guideline) checks exist only for the crops listed in
    outputs/treatment_timing_crosswalk.csv.

Run from the repository root:  RAW_DIR=/path/to/raw python audit_claude_20260930/scripts/02_full_panel_check.py
Writes outputs/full_panel_check_cells.csv and outputs/full_panel_check_log.txt.
"""
import datetime
import hashlib
import os
import platform
import re
import sys

import warnings

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
sys.path.insert(0, os.path.dirname(__file__))
from xml2003 import read as read_xml  # noqa: E402

RAW_DIR = os.environ.get('RAW_DIR', '/mnt/user-data/uploads/')
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(ROOT, 'audit_claude_20260930', 'outputs')
FILES = ['01_fruit_total.xls', '02_fruit_bearing.xlsx', '03_rice.xlsx', '04_beans.xlsx', '05_coarse_grains.xlsx',
         '06_potatoes.xlsx', '07_barely.xlsx', '08_vegetables_root.xls', '09_vegetables_leafy.xls',
         '10_vegetables_spices.xls', '13_특용작물생산량__땅콩_.xlsx']
NATIONAL = {'계', '전국'}


def num(v):
    """Panel rule: non-numeric ('-', blank) and non-positive national values are treated as missing."""
    try:
        x = float(str(v).replace(',', '').strip())
    except (TypeError, ValueError):
        return np.nan
    return x if x > 0 else np.nan


def split_item(s):
    m = re.match(r'(.*):(.*?)\s*[\(\[]([^\)\]]*)[\)\]]\s*$', str(s).strip())
    return (m.group(1).strip(), m.group(2).strip()) if m else (str(s).strip(), '')


cells, meta_log = {}, []
for f in FILES:
    path = os.path.join(RAW_DIR, f)
    sha = hashlib.sha256(open(path, 'rb').read()).hexdigest()
    if open(path, 'rb').read(200).lstrip().startswith(b'<?xml'):
        d = read_xml(path)
        tid = next(str(r[1]).strip() for r in d['메타정보'] if r and len(r) > 1 and r[0] and '통계표ID' in str(r[0]))
        rows = d['데이터']
        yc = [(i, int(str(h)[:4])) for i, h in enumerate(rows[1]) if h and re.match(r'\d{4}', str(h))]
        reg = None
        for r in rows[2:]:
            if r[0]:
                reg = r[0].strip()
            if reg not in NATIONAL or not r[1]:
                continue
            crop, var = split_item(r[1])
            for i, y in yc:
                cells[(tid, crop, var, y)] = (r[i], f)
        reader = 'xml2003 (same parser as build)'
    else:
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)  # read_only mode truncates rows in these exports
        meta = [row for row in wb['메타정보'].iter_rows(values_only=True)]
        tid = next(str(r[1]).strip() for r in meta if r and len(r) > 1 and r[0] and '통계표ID' in str(r[0]))
        rows = [row for row in wb['데이터'].iter_rows(values_only=True)]
        two_level = str(rows[0][1]).startswith('시도별')
        start = 2 if two_level else 1
        reg1 = None
        for r in rows[2:]:
            if r[0]:
                reg1 = str(r[0]).strip()
            r2 = str(r[1]).strip() if two_level else ''
            if reg1 not in NATIONAL or (two_level and r2 != '소계'):
                continue
            yr = None
            for ci in range(start, len(rows[0])):
                if rows[0][ci] is not None:  # merged year headers: carry forward
                    yr = int(str(rows[0][ci])[:4])
                crop, var = split_item(rows[1][ci])
                cells[(tid, crop, var, yr)] = (r[ci], f)
        reader = 'openpyxl (independent of build reader)'
    meta_log.append(f'{f}\t{tid}\tSHA-256={sha}\treader={reader}')

V = pd.read_csv(os.path.join(ROOT, 'verified_results', 'verified_master_panel.csv'))
SPEC = [('area_ha', '면적', 0), ('yield_kg10a', '10a당 생산량', 0), ('production_t', '생산량', None),  # production: all crops
        ('total_orchard_area_ha', '면적', 1), ('bearing_area_ha', '면적', 1)]
out = []
for r in V.itertuples():
    lab = r.kosis_production_label_original
    for col, var, per in SPEC:
        if per is not None and r.perennial != per:
            continue
        tid = 'DT_1ET0296' if col == 'bearing_area_ha' else r.table_id
        crop_label = '논벼:조곡' if (r.crop_id == 'rice' and col == 'production_t') else lab
        key = (tid, crop_label, var, int(r.year))
        raw, src = cells.get(key, (None, ''))
        rv = num(raw)
        pv = getattr(r, col)
        status = ('both_missing' if np.isnan(rv) and pd.isna(pv) else
                  'missing_mismatch' if np.isnan(rv) != pd.isna(pv) else
                  'exact' if rv == pv else 'value_mismatch')
        out.append(dict(crop_id=r.crop_id, year=r.year, variable=col, table_id=tid, kosis_label=crop_label,
                        source_file=src, raw=raw, raw_numeric=rv, panel=pv,
                        abs_diff=abs(rv - pv) if status in ('exact', 'value_mismatch') else np.nan, status=status))
C = pd.DataFrame(out)
C.to_csv(os.path.join(OUT, 'full_panel_check_cells.csv'), index=False, encoding='utf-8-sig')

TC = pd.read_csv(os.path.join(ROOT, 'verified_results', 'verified_treatment_coding.csv'))
tv = V.drop_duplicates('crop_id').set_index('crop_id')[['pilot_year', 'national_year']]
tm = TC.set_index('crop_id')[['pilot_year', 'national_year']]
common = tv.index.intersection(tm.index)
treat_mismatch = [(c, tuple(tv.loc[c]), tuple(tm.loc[c])) for c in common
                  if not all((pd.isna(a) and pd.isna(b)) or a == b for a, b in zip(tv.loc[c], tm.loc[c]))]

with open(os.path.join(OUT, 'full_panel_check_log.txt'), 'w', encoding='utf-8') as fh:
    fh.write(f'run_utc={datetime.datetime.now(datetime.timezone.utc).isoformat()}\npython={platform.python_version()} pandas={pd.__version__} '
             f'openpyxl={openpyxl.__version__} numpy={np.__version__}\nRAW_DIR={RAW_DIR}\n\n[input files]\n')
    fh.write('\n'.join(meta_log) + '\n\n')
    fh.write(f'verified panel rows={len(V)} crops={V.crop_id.nunique()} years={V.year.min()}-{V.year.max()}\n')
    fh.write(f'numeric cells compared={len(C)}\n')
    fh.write(C.groupby(['variable', 'status']).size().unstack(fill_value=0).to_string() + '\n\n')
    fh.write(f'max abs diff among non-missing cells={C.abs_diff.max()}\n')
    fh.write(f'value_mismatch={int((C.status == "value_mismatch").sum())} missing_mismatch={int((C.status == "missing_mismatch").sum())}\n')
    nf = C[C.source_file == '']
    fh.write(f'raw cells not found={len(nf)} (all have panel value missing: {bool(nf.panel.isna().all())})\n')
    fh.write('not-found cells by variable and crop (year range):\n')
    fh.write(nf.groupby(['variable', 'crop_id']).year.agg(['min', 'max', 'count']).to_string() + '\n\n[treatment years: internal consistency only]\n')
    fh.write(f'crops compared={len(common)} mismatches={len(treat_mismatch)} {treat_mismatch}\n')
print(open(os.path.join(OUT, 'full_panel_check_log.txt'), encoding='utf-8').read())
